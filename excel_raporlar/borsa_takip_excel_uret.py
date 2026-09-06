"""TÜRİB + TMO + 4 Ticaret Borsası (Edirne/Bandırma/Kırklareli/Tekirdağ) fiyat
takip Excel'ini üretir - sadece buğday/arpa/mısır.

Kullanım:
    python3 -m excel_raporlar.borsa_takip_excel_uret [cikti_yolu]

Varsayilan cikti: ~/Desktop/borsa_takip_arpa_bugday_misir.xlsx

TASARIM KURALLARI (kullanıcının 2026-09-05'te elle düzenleyip "bundan sonra
hep böyle" dediği şablon - values are learned from the user's manual Numbers edit):
- Her "veri" sayfasında (referans/açıklama sayfaları haric) AutoFilter acik
- Baslik satiri: Verdana, bold, ortali, ince kenarlik, dolgu renkli
- Veri satirlari: Verdana, 10pt
- Fiyat sutunlari: Finansal (Muhasebe) bicimi, ₺ sembolu ONEK olarak (TL yazisi degil)
- Tutar (TL) sutunu: Finansal bicim + "TL" son eki (buyuk TL tutarlari icin)
- Miktar/Hacim/Islem Sayisi: Finansal (Muhasebe) tamsayi bicimi (ondalik yok)
- Tarih sutunu: GERCEK tarih tipi (metin degil) - dd.mm.yyyy gorunumlu,
  boylece kronolojik siralama/filtre dogru calisir
- Sutun genislikleri icerige gore otomatik ayarlanir
"""
import argparse
import json
import sqlite3
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

REPO_KOKU = Path(__file__).parent.parent
DB_PATH = REPO_KOKU / "veri_kaynagi" / "borsa_verileri.db"
TURIB_GUNLUK_VERI_DIZINI = REPO_KOKU / "turib_2023_2026_tam" / "gunluk_veri"
VARSAYILAN_CIKTI = Path.home() / "Desktop" / "borsa_takip_arpa_bugday_misir.xlsx"


# ============================================================================
# Yardimci fonksiyonlar
# ============================================================================

def ilgili_urun(urun: str) -> bool:
    """Sadece bugday/arpa/misir (ve dogrudan turevleri) - kanola/aycicek/saman vb disarida."""
    u = (urun or "").upper()
    return ("BUĞDAY" in u) or u.startswith("ARPA") or ("MISIR" in u)


def parse_tl(deger) -> float | None:
    """Karisik bicimli TL tutarlarini (Bandirma US-stili '66,850,550.66',
    TDAG/Kirklareli/ETB TR-stili '660.080,00') dogru sayiya cevirir. Kural:
    son gecen ',' veya '.' hangisiyse o ondalik ayracidir, digeri binlik
    ayraci sayilip silinir."""
    if deger is None:
        return None
    s = str(deger).strip()
    if s == "" or s.lower() == "nan":
        return None
    if s.rfind(",") > s.rfind("."):
        s2 = s.replace(".", "").replace(",", ".")
    else:
        s2 = s.replace(",", "")
    try:
        return float(s2)
    except ValueError:
        return None


def parse_int(deger) -> int | None:
    if deger is None:
        return None
    s = str(deger).strip()
    if s == "" or s.lower() == "nan":
        return None
    try:
        return int(float(s))
    except (ValueError, TypeError):
        return None


def tr_sayi(s):
    if s is None or s == "":
        return None
    return float(str(s).replace(".", "").replace(",", "."))


# ============================================================================
# Satis sekli / islem turu aciklamalari (kaynak: borsa.tobb.org.tr/islem_turu.php
# + etb.org.tr/aylikbulten/aylik-bulten sayfasindaki resmi legend + gozlem/tahmin)
# ============================================================================

SATIS_SEKLI_ACIKLAMA = {
    "FABA": ("Fabrika alım", "kesin (TOBB resmi liste)"),
    "FABS": ("Fabrika satım", "kesin (TOBB resmi liste)"),
    "HMS": ("Hazır müstahsil satış", "kesin (TOBB resmi liste)"),
    "HTS": ("Hazır tüccar satış", "kesin (TOBB resmi liste)"),
    "KOOPA": ("Kooperatif alım", "kesin (TOBB resmi liste)"),
    "KOOPS": ("Kooperatif satım", "kesin (TOBB resmi liste)"),
    "TARA": ("Tariş alım", "kesin (TOBB resmi liste)"),
    "TARS": ("Tariş satım", "kesin (TOBB resmi liste)"),
    "TMOA": ("Toprak Mahsulleri Ofisi alım", "kesin (TOBB resmi liste)"),
    "TMOS": ("Toprak Mahsulleri Ofisi satım", "kesin (TOBB resmi liste)"),
    "VMS": ("Vadeli müstahsil satış", "kesin (TOBB resmi liste)"),
    "VTS": ("Vadeli tüccar satış", "kesin (TOBB resmi liste)"),
    "H.M.S": ("Hazır Müst.Satışı", "kesin (ETB resmi liste)"),
    "H.T.S": ("Hazır Tüccar Satışı", "kesin (ETB resmi liste)"),
    "TMO.A": ("Toprak Mahsulleri Ofisi Alış", "kesin (ETB resmi liste)"),
    "TMO.S": ("Toprak Mahsulleri Ofisi Satış", "kesin (ETB resmi liste)"),
    "DÜÇ.S": ("Devlet Üretme Çiftliği Satış", "kesin (ETB resmi liste)"),
    "VHMS": ("Vadeli Hz.Müs.Satışı", "kesin (ETB resmi liste)"),
    "İ.T.A": ("İthal Tüccar Alışı", "kesin (ETB resmi liste)"),
    "ŞFBA": ("Şeker Fabrikası Alışı", "kesin (ETB resmi liste)"),
    "TMOİKA": ("TMO İhr. Kayıt Satışı", "kesin (ETB resmi liste)"),
    "TMOAL": ("TMO Alivre Alışı", "kesin (ETB resmi liste)"),
    "VTMOA": ("Vadeli TMO Alışı", "kesin (ETB resmi liste)"),
    "VTMOS": ("Vadeli TMO Satışı", "kesin (ETB resmi liste)"),
    "THKST": ("T.H.K Satışı", "kesin (ETB resmi liste)"),
    "İ.K.S": ("İhraç Kayıtlı Satış", "kesin (ETB resmi liste)"),
    "KHMS": ("Kaza Hz.Müs.Satışı", "kesin (ETB resmi liste)"),
    "KHTS": ("Kaza Hz. Tüccar Satışı", "kesin (ETB resmi liste)"),
    "KTMOA": ("Kaza T.M.O.Alışı / Kz.TMO.Alivre Alışı", "kesin (ETB resmi liste, kod ETB legend'inde 2 kez farklı acilimla gecti)"),
    "KTMOS": ("Kaza T.M.O Satışı", "kesin (ETB resmi liste)"),
    "KDÜÇS": ("Kaza Düç. Satış", "kesin (ETB resmi liste)"),
    "KVHMS": ("K.Vd.Hz.Müs.Satışı", "kesin (ETB resmi liste)"),
    "KVHTS": ("K.VD.Tüccar Satışı", "kesin (ETB resmi liste)"),
    "KİTA": ("Kz.İthal Tüccar Alım", "kesin (ETB resmi liste)"),
    "KŞFBA": ("Kz.Şeker Fb.Alışı", "kesin (ETB resmi liste)"),
    "KDİSK": ("Kz.TMO.İhr.Kayıt Satış", "kesin (ETB resmi liste)"),
    "KVDAL": ("Kz.Vd.TMO.Alışı", "kesin (ETB resmi liste)"),
    "KVDST": ("Kz.Vd.TMO.Satışı", "kesin (ETB resmi liste)"),
    "TS": ("Tüccar satış (HTS'nin kısa varyantı olabilir) - TEYIT EDILEMEDI, kullanıcının 'teminatlı satış' tahmini de olası", "TAHMIN"),
    "GMS": ("Geç müstahsil satış (olası) - TEYIT EDILEMEDI", "TAHMIN"),
    "KOP.A": ("Kooperatif alım (KOOPA'nın borsaya özgü kısaltması)", "TAHMIN (yüksek güven)"),
    "KOP.S": ("Kooperatif satım (KOOPS'un borsaya özgü kısaltması)", "TAHMIN (yüksek güven)"),
    "Ü.ÇKS.SAT": ("Üretici ÇKS (Çiftçi Kayıt Sistemi) satış", "TAHMIN (yüksek güven)"),
    "Y.T.K.AL.": ("Yetkili Tüccar/Kurum alım", "TAHMIN"),
    "Y.T.K.SAT.": ("Yetkili Tüccar/Kurum satış", "TAHMIN"),
    "OF.A.": ("Ofis alım (TMO'nun 'Ofis' olarak anıldığı alım - TMOA'ya karşılık gelebilir)", "TAHMIN (yüksek güven)"),
    "İHRC.": ("İhracat satışı", "TAHMIN (yüksek güven)"),
    "İTHAL": ("İthalat alışı", "TAHMIN (yüksek güven)"),
    "GTM.S": ("Tanımsız - TEYIT EDILEMEDI", "TAHMIN (düşük güven)"),
    "H:T:S": ("Hazır tüccar satış (HTS eşdeğeri, Bandırma'nın kendi formatı)", "TAHMIN (yüksek güven)"),
    "S:M:S": ("Tanımsız (muhtemelen bir satış türü kombinasyonu) - TEYIT EDILEMEDI", "TAHMIN (düşük güven)"),
    "V.SAT": ("Vadeli satış (VMS/VTS ailesi)", "TAHMIN (orta güven)"),
    "OF.S.": ("Ofis satış (OF.A.'nın satış eşdeğeri, TMOS'a karşılık gelebilir)", "TAHMIN (yüksek güven)"),
    "KOOP.A": ("Kooperatif alım (KOOPA eşdeğeri)", "TAHMIN (yüksek güven)"),
    "H:T:A": ("Hazır tüccar alışı (H:T:S'nin alım eşdeğeri)", "TAHMIN (yüksek güven)"),
    "TMO.AL": ("TMO alış (TMOA eşdeğeri)", "TAHMIN (yüksek güven)"),
    "TMO.ST.": ("TMO satış (TMOS eşdeğeri)", "TAHMIN (yüksek güven)"),
    "TMOST": ("TMO satış (TMOS eşdeğeri)", "TAHMIN (yüksek güven)"),
    "İTHLT": ("İthalat (İTHAL eşdeğeri)", "TAHMIN (yüksek güven)"),
    "V:ALŞ": ("Vadeli alış", "TAHMIN (orta güven)"),
    "TBOA": ("Tanımsız (belki 'Ticaret Borsası Ofis Alım') - TEYIT EDILEMEDI", "TAHMIN (düşük güven)"),
    "TBOS": ("Tanımsız (belki 'Ticaret Borsası Ofis Satış') - TEYIT EDILEMEDI", "TAHMIN (düşük güven)"),
    "GTTS": ("Tanımsız - TEYIT EDILEMEDI", "TAHMIN (düşük güven)"),
    "G.F.U": ("Tanımsız - TEYIT EDILEMEDI", "TAHMIN (düşük güven)"),
}


# ============================================================================
# Hucre bicimleri (kullanicinin elle sectigi "Finansal" bicimlerle birebir)
# ============================================================================

PARA_TRY_FORMAT = '_-"₺"* #,##0.00_-;-"₺"* #,##0.00_-;_-"₺"* "-"??_-;_-@_-'
PARA_TUTAR_TL_FORMAT = '_-* #,##0.00\\ "TL"_-;-* #,##0.00\\ "TL"_-;_-* "-"??\\ "TL"_-;_-@_-'
PARA_USD_FORMAT = '_-$* #,##0.00_-;-$* #,##0.00_-;_-$* "-"??_-;_-@_-'
TAMSAYI_FORMAT = '_-* #,##0_-;-* #,##0_-;_-* "-"_-;_-@_-'
TARIH_FORMAT = "dd.mm.yyyy"
# Deger zaten yuzde sayisi olarak saklaniyor (4.92 = %4,92) - Excel'in yerlesik
# '%' bicimi degeri 100 ile carpar, o yuzden duz metin eki kullaniliyor.
YUZDE_FORMAT = '0.00"%"'

BASLIK_FONT = Font(name="Verdana", size=10, bold=True)
BASLIK_FILL = PatternFill("solid", fgColor="4472C4")
BASLIK_ALIGN = Alignment(horizontal="center", vertical="center")
INCE_KENAR = Side(style="thin", color="B0B0B0")
BASLIK_BORDER = Border(left=INCE_KENAR, right=INCE_KENAR, top=INCE_KENAR, bottom=INCE_KENAR)
VERI_FONT = Font(name="Verdana", size=10)


def sayfayi_bicimlendir(ws, columns, tablo_adi: str, autofilter=True):
    """Basligi biciimlendirir, sutun bazinda finansal/tarih/tamsayi formati
    uygular, sutun genisligini icerige gore ayarlar, AutoFilter ekler."""
    max_row = ws.max_row
    max_col = len(columns)

    for idx, col_adi in enumerate(columns, start=1):
        col_letter = get_column_letter(idx)
        col_upper = col_adi.upper()

        # baslik hucresi
        hucre = ws[f"{col_letter}1"]
        hucre.font = BASLIK_FONT
        hucre.fill = BASLIK_FILL
        hucre.alignment = BASLIK_ALIGN
        hucre.border = BASLIK_BORDER

        if col_upper == "TARIH":
            fmt = TARIH_FORMAT
        elif "%" in col_adi:
            # '#,##0.00" TL"' gibi para bicimleri yuzde sutununa (Net Marj (%)
            # gibi) yanlislikla uygulanmasin diye para kontrollerinden ONCE
            # kontrol edilir - deger zaten yuzde sayisi (4.92 = %4,92), Excel'in
            # carpma yapan yerlesik '%' bicimi degil, duz metin eki kullanilir.
            fmt = YUZDE_FORMAT
        elif "USD" in col_upper:
            fmt = PARA_USD_FORMAT
        elif col_upper.startswith("TUTAR"):
            fmt = PARA_TUTAR_TL_FORMAT
        elif "FIYAT" in col_upper or "MARJ" in col_upper:
            fmt = PARA_TRY_FORMAT
        elif "MIKTAR" in col_upper or "HACIM" in col_upper or "İŞLEM SAYISI" in col_upper:
            fmt = TAMSAYI_FORMAT
        else:
            fmt = None

        # sutun genisligi: baslik + ornek veri uzunluguna gore
        ornek_uzunluk = max(
            [len(str(ws.cell(row=r, column=idx).value or "")) for r in range(2, min(max_row, 50) + 1)],
            default=0,
        )
        ws.column_dimensions[col_letter].width = min(max(len(col_adi), ornek_uzunluk) + 3, 70)

        if fmt:
            for row in range(2, max_row + 1):
                ws[f"{col_letter}{row}"].number_format = fmt

        for row in range(2, max_row + 1):
            ws[f"{col_letter}{row}"].font = VERI_FONT

    if autofilter and max_row > 1:
        son_sutun = get_column_letter(max_col)
        ws.auto_filter.ref = f"A1:{son_sutun}{max_row}"
    ws.freeze_panes = "A2"


def aciklama_ekle(df, satis_sekli_kolonu):
    df["satis_sekli_aciklama"] = df[satis_sekli_kolonu].map(
        lambda x: SATIS_SEKLI_ACIKLAMA.get(x, ("Bilinmiyor", "yok"))[0]
    )
    df["satis_sekli_guvenilirlik"] = df[satis_sekli_kolonu].map(
        lambda x: SATIS_SEKLI_ACIKLAMA.get(x, ("Bilinmiyor", "yok"))[1]
    )
    return df


# ============================================================================
# TÜRİB gunluk agregasyon (buğday/arpa/mısır, tum siniflar)
# ============================================================================

def turib_gunluk_urun_grubu(cls: str) -> str | None:
    c = cls.upper()
    if "BUĞDAY" in c:
        return "BUĞDAY"
    if "ARPA" in c:
        return "ARPA"
    if c.startswith("MISIR"):
        return "MISIR"
    return None


def turib_dataframe_uret() -> pd.DataFrame:
    import glob
    rows = []
    for f in sorted(glob.glob(str(TURIB_GUNLUK_VERI_DIZINI / "*.json"))):
        d = json.load(open(f))
        tarih = d["tarih"]
        for section in ("normal_seans", "anlasmali"):
            for row in d.get(section, []):
                cls = row.get("Enstrüman Sınıfı", "")
                grup = turib_gunluk_urun_grubu(cls)
                if grup is None:
                    continue
                fiyat = tr_sayi(row.get("Ağırlıklı Ortalama Fiyat")) or tr_sayi(row.get("Kapanış Fiyatı"))
                miktar = tr_sayi(row.get("İşlem Miktarı [KG]"))
                rows.append({"tarih": tarih, "grup": grup, "sinif": cls, "fiyat_tl_kg": fiyat, "miktar_kg": miktar})

    df = pd.DataFrame(rows)

    def agg(g):
        fiyatli = g.dropna(subset=["fiyat_tl_kg"])
        if len(fiyatli) and fiyatli["miktar_kg"].sum() > 0:
            agirlikli_ort = np.average(fiyatli["fiyat_tl_kg"], weights=fiyatli["miktar_kg"])
            min_f, max_f = fiyatli["fiyat_tl_kg"].min(), fiyatli["fiyat_tl_kg"].max()
        else:
            agirlikli_ort = min_f = max_f = None
        return pd.Series({
            "Toplam Hacim (KG)": g["miktar_kg"].sum(skipna=True),
            "İşlem Sayısı (satır)": len(g),
            "Min Fiyat (TL/ton)": min_f * 1000 if min_f else None,
            "Ağırlıklı Ort. Fiyat (TL/ton)": agirlikli_ort * 1000 if agirlikli_ort else None,
            "Max Fiyat (TL/ton)": max_f * 1000 if max_f else None,
        })

    daily = df.groupby(["tarih", "grup", "sinif"]).apply(agg, include_groups=False).reset_index()
    daily = daily.rename(columns={"tarih": "Tarih", "grup": "Ürün Grubu", "sinif": "Enstrüman Sınıfı"})
    daily["Tarih"] = pd.to_datetime(daily["Tarih"]).dt.date
    return daily.sort_values(["Ürün Grubu", "Enstrüman Sınıfı", "Tarih"])


# ============================================================================
# TMO
# ============================================================================

def tmo_dataframe_uret(conn) -> pd.DataFrame:
    rows = []
    for row in conn.execute("SELECT * FROM fiyatlar WHERE kaynak='TMO' ORDER BY tarih, il, urun"):
        if not ilgili_urun(row["urun"]):
            continue
        ham = json.loads(row["ham_veri"]) if row["ham_veri"] else {}
        rows.append({
            "Tarih": row["tarih"], "İl/Borsa": row["il"], "Ürün": row["urun"],
            "Miktar (ton)": ham.get("miktar_ton"),
            "Fiyat (TL/ton)": ham.get("tl_ton"), "Fiyat (USD/ton)": ham.get("usd_ton"),
            "Önceki Dönem Miktar (ton)": ham.get("onceki_donem_miktar_ton"),
            "Önceki Dönem Fiyat (TL/ton)": ham.get("onceki_donem_tl_ton"),
            "Önceki Dönem Fiyat (USD/ton)": ham.get("onceki_donem_usd_ton"),
            "Geçen Yıl Fiyat (TL/ton)": ham.get("gecen_yil_tl_ton"),
        })
    df = pd.DataFrame(rows)
    if len(df):
        df["Tarih"] = pd.to_datetime(df["Tarih"]).dt.date
    return df


# ============================================================================
# Ticaret borsalari (Edirne/Bandirma/Kirklareli/Tekirdag)
# ============================================================================

def borsa_dataframe_uret(conn, kaynak_kodlari: list[str]) -> pd.DataFrame:
    rows = []
    for kaynak_kod in kaynak_kodlari:
        for row in conn.execute("SELECT * FROM fiyatlar WHERE kaynak=? ORDER BY tarih, urun", (kaynak_kod,)):
            if not ilgili_urun(row["urun"]):
                continue
            ham = json.loads(row["ham_veri"]) if row["ham_veri"] else {}
            satis_sekli = ham.get("satis_sekli") or ham.get("SatisSekliKodu") or ham.get("satisSekliKodu")
            islem_sayisi = (ham.get("islem_sayisi") or ham.get("IslemAdedi")
                             or ham.get("islemSayisi") or ham.get("İşlem Adedi"))
            rows.append({
                "Kaynak": kaynak_kod, "Tarih": row["tarih"], "İl": row["il"], "İlçe": row["ilce"],
                "Ürün": row["urun"], "Detay/Kod": row["detay"],
                "Min Fiyat (TL/kg)": row["min_fiyat"], "Ort. Fiyat (TL/kg)": row["ort_fiyat"],
                "Max Fiyat (TL/kg)": row["max_fiyat"],
                "Miktar (KG)": row["miktar"], "Birim": row["birim"],
                "Satış Şekli Kodu": satis_sekli,
                "İşlem Sayısı": parse_int(islem_sayisi),
                "Tutar (TL)": parse_tl(ham.get("tutar_tl") or ham.get("tutar")),
            })
    df = pd.DataFrame(rows)
    if len(df):
        df["Tarih"] = pd.to_datetime(df["Tarih"]).dt.date
        if "Satış Şekli Kodu" in df.columns and df["Satış Şekli Kodu"].notna().any():
            df = aciklama_ekle(df, "Satış Şekli Kodu")
    return df


# ============================================================================
# Ana
# ============================================================================

def uret(cikti_yolu: Path):
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row

    turib_df = turib_gunluk_veri_var_mi_kontrol_et_ve_uret()
    tmo_df = tmo_dataframe_uret(conn)

    borsa_gruplari = [
        (["ETB", "ETB_AYLIK"], "Edirne Ticaret Borsası"),
        (["BANDIRMA"], "Bandırma Ticaret Borsası"),
        (["KIRKLARELI_AYLIK"], "Kırklareli Ticaret Borsası"),
        (["TDAG"], "Tekirdağ Ticaret Borsası"),
    ]
    borsa_dfleri = [(ad, borsa_dataframe_uret(conn, kodlar)) for kodlar, ad in borsa_gruplari]
    conn.close()

    aciklama_df = pd.DataFrame([
        {"Kısaltma": k, "Açıklama": v[0], "Güvenilirlik": v[1]}
        for k, v in SATIS_SEKLI_ACIKLAMA.items()
    ])

    with pd.ExcelWriter(cikti_yolu, engine="openpyxl") as writer:
        turib_df.to_excel(writer, sheet_name="TÜRİB", index=False)
        sayfayi_bicimlendir(writer.sheets["TÜRİB"], turib_df.columns, "TÜRİB")

        tmo_df.to_excel(writer, sheet_name="TMO", index=False)
        sayfayi_bicimlendir(writer.sheets["TMO"], tmo_df.columns, "TMO")

        for sheet_adi, df in borsa_dfleri:
            ad = sheet_adi[:31]
            df.to_excel(writer, sheet_name=ad, index=False)
            sayfayi_bicimlendir(writer.sheets[ad], df.columns, ad)

        aciklama_df.to_excel(writer, sheet_name="Satış Şekli Açıklamaları", index=False)
        sayfayi_bicimlendir(writer.sheets["Satış Şekli Açıklamaları"], aciklama_df.columns,
                             "Satış Şekli Açıklamaları", autofilter=False)

    print(f"Excel yazildi: {cikti_yolu}")
    print(f"TÜRİB: {len(turib_df)} satir | TMO: {len(tmo_df)} satir")
    for ad, df in borsa_dfleri:
        print(f"{ad}: {len(df)} satir")


def turib_gunluk_veri_var_mi_kontrol_et_ve_uret() -> pd.DataFrame:
    if not TURIB_GUNLUK_VERI_DIZINI.exists():
        raise FileNotFoundError(f"TÜRİB gunluk veri dizini bulunamadi: {TURIB_GUNLUK_VERI_DIZINI}")
    return turib_dataframe_uret()


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("cikti", nargs="?", default=str(VARSAYILAN_CIKTI))
    args = parser.parse_args()
    uret(Path(args.cikti))
