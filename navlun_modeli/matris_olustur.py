"""navlun_formulu.py'deki modeli km x ton matrisi olarak Excel'e doker (sabit/degisken/toplam ayri)."""
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill
from navlun_formulu import navlun_hesapla

KM_LISTESI = [15, 25, 40, 55, 85, 100, 140, 150, 190, 210, 300, 400, 500, 700]
TON_LISTESI = [0.5, 1, 2, 5, 10, 15, 20, 27, 40]

KOK = Path(__file__).parent
wb = openpyxl.Workbook()

# --- Sayfa 1: TL / SEVKIYAT (toplam navlun faturasi) ---
ws1 = wb.active
ws1.title = "TL-sevkiyat (toplam)"
ws1.cell(1, 1, "km \\ ton").font = Font(bold=True)
for j, ton in enumerate(TON_LISTESI, start=2):
    ws1.cell(1, j, ton).font = Font(bold=True)
for i, km in enumerate(KM_LISTESI, start=2):
    ws1.cell(i, 1, km).font = Font(bold=True)
    for j, ton in enumerate(TON_LISTESI, start=2):
        r = navlun_hesapla(km, ton)
        ws1.cell(i, j, r["toplam_tl_sevkiyat"])

# --- Sayfa 2: TL / TON (arbitraj karsilastirmasi icin asil kullanilan) ---
ws2 = wb.create_sheet("TL-ton (toplam)")
ws2.cell(1, 1, "km \\ ton").font = Font(bold=True)
for j, ton in enumerate(TON_LISTESI, start=2):
    ws2.cell(1, j, ton).font = Font(bold=True)
for i, km in enumerate(KM_LISTESI, start=2):
    ws2.cell(i, 1, km).font = Font(bold=True)
    for j, ton in enumerate(TON_LISTESI, start=2):
        r = navlun_hesapla(km, ton)
        c = ws2.cell(i, j, r["toplam_tl_ton"])
        if r["guven"] == "interpolasyon":
            c.fill = PatternFill("solid", fgColor="FFF2CC")  # sari = tahmini/interpolasyon

# --- Sayfa 3: Detay tablo (sabit + degisken ayri, guven etiketli) ---
ws3 = wb.create_sheet("Detay (sabit+degisken)")
basliklar = ["km", "ton", "yuk_bandi", "guven", "arac_sayisi",
             "sabit_tl_sevkiyat", "degisken_tl_sevkiyat", "toplam_tl_sevkiyat",
             "sabit_tl_ton", "degisken_tl_ton", "toplam_tl_ton"]
for j, b in enumerate(basliklar, start=1):
    ws3.cell(1, j, b).font = Font(bold=True)
row = 2
for km in KM_LISTESI:
    for ton in TON_LISTESI:
        r = navlun_hesapla(km, ton)
        for j, b in enumerate(basliklar, start=1):
            ws3.cell(row, j, r[b])
        row += 1

# --- Sayfa 4: Aciklama / kaynaklar ---
ws4 = wb.create_sheet("Aciklama")
aciklama = [
    "GENEL NAVLUN FORMULU - km VE ton hassasiyetli (2026-09-05)",
    "",
    "Her hucre SABIT + DEGISKEN maliyetin toplami (breakdown 'Detay' sayfasinda ayri kolonlarda).",
    "",
    "3 yuk bandi:",
    "1) Parsiyel/LCL (<=2 ton): sabit=4.000 TL + K(km)*km  [K: 0-100km->95, 100-250km->65, 250km+->52]",
    "   Kaynak: ~/Desktop/agents/skill.md (NORA Global, n~9 gercek teklif, Haziran 2026)",
    "2) Tam yuk/Dorse (20-27 ton, dokme hububat): sabit=4.416 TL + 81,79 TL/km",
    "   Kaynak: 9 referans (7 skill.md konteyner/liman drenaji + Cem Yoluk/NORA Global 2 mail: "
    "Ankara-Izmit 30.450 TL, Aliaga-ESBAS 19.500 TL [outlier, regresyon disi])",
    "3) Orta olcek/kismi dorse (2-20 ton): GERCEK VERI YOK, Bant1<->Bant2 arasi tonaja gore "
    "DOGRUSAL INTERPOLASYON. 'TL-ton' sayfasinda sari renkle isaretli - bunlar TAHMINIDIR.",
    "",
    "40 ton+ orneklerinde arac_sayisi>1 (birden fazla dorse gerekiyor, maliyet orantili buyutulur).",
    "",
    "Onemli: Bu HALA kaba bir modeldir (±%15-30 sapma normal), ozellikle Bant 3 (interpolasyon).",
    "Yeni gercek teklif (ozellikle 2-20 ton araliginda) geldikce navlun_formulu.py guncellenmeli.",
    "",
    "Kod: navlun_formulu.py (fonksiyon: navlun_hesapla(km, ton))",
]
for i, line in enumerate(aciklama, start=1):
    ws4.cell(i, 1, line)
ws4.column_dimensions["A"].width = 110

out = KOK / "navlun_matrisi.xlsx"
wb.save(out)
print(f"Kaydedildi: {out}")
