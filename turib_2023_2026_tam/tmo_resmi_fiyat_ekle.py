"""Son TMO ile ilgili konusmalardan cikan resmi alis/satis fiyati verisini
mevcut tmo_turib_bu_hafta.xlsx dosyasina yeni bir sekme olarak ekler."""
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

HEDEF = Path("/Users/utkus/Desktop/tmo_turib_bu_hafta.xlsx")

wb = openpyxl.load_workbook(HEDEF)
if "TMO 2026 Resmi Fiyatlar" in wb.sheetnames:
    del wb["TMO 2026 Resmi Fiyatlar"]
ws = wb.create_sheet("TMO 2026 Resmi Fiyatlar")

basliklar = ["ürün", "kalem", "tl_ton", "açıklama"]
ws.append(basliklar)
for c in ws[1]:
    c.font = Font(bold=True)

satirlar = [
    ("Arpa", "TMO Alım Fiyatı (taban)", 12750, "Çiftçiden alınan resmi taban fiyat"),
    ("Arpa", "+ Devlet Desteği ile Üreticiye Geçen", 15764, "Alım fiyatı + destekleme ödemesi"),
    ("Arpa", "TMO Satış Fiyatı", 14000, "TMO'nun 1 Ekim 2026'dan itibaren sattığı fiyat"),
    ("Arpa", "TMO Brüt Marj (Satış - Alım)", 1250, "Destekler hariç"),
    ("Buğday (Ekmeklik/Makarnalık)", "TMO Alım Fiyatı (taban)", 16500, "Çiftçiden alınan resmi taban fiyat"),
    ("Buğday (Ekmeklik/Makarnalık)", "+ Devlet Desteği ile Üreticiye Geçen", 19514, "Alım fiyatı + destekleme ödemesi"),
    ("Buğday (Ekmeklik/Makarnalık)", "TMO Satış Fiyatı", 18500, "TMO'nun sattığı fiyat"),
    ("Buğday (Ekmeklik/Makarnalık)", "TMO Brüt Marj (Satış - Alım)", 2000, "Destekler hariç"),
]
for row in satirlar:
    ws.append(row)

ws.append([])
ws.append(["Kaynak", "T.C. Tarım ve Orman Bakanlığı resmi web sitesi", "", ""])
ws.append(["Açıklama tarihi", "2 Haziran 2026", "", ""])
ws.append(["URL", "https://www.tarimorman.gov.tr/Haber/7092/2026-Yili-Hububat-Alim-Ve-Satis-Fiyatlari-Belirlendi", "", ""])
ws.append([])
ws.append(["Not", "Bu haftaki (1-4 Eylül 2026) tek TMO arpa işlemi (ISIN TRXTTDAA2610, "
                  "Kırşehir/Mucur, 55,72 ton) 'Anlaşmalı' türde - TÜRİB anlaşmalı işlemlerde "
                  "fiyat yayınlamıyor, bu yüzden bu işlemin gerçek fiyatı bilinmiyor. "
                  "'ISIN Özet (hafta)' sekmesine bakınız.", "", ""])

for i, w in enumerate((32, 38, 12, 60), start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

wb.save(HEDEF)
print(f"Kaydedildi: {HEDEF} (sekmeler: {wb.sheetnames})")
