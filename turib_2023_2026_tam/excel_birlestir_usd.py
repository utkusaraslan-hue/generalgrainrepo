"""TURIB 2023-2026 gunluk bultenlerini, o GUNUN TCMB USD Doviz Satis kuruyla
TL degerlerini USD'ye cevirerek tek Excel'de birlestirir (excel_birlestir.py'nin
USD versiyonu - ayni yapi + her satirda kullanilan kur seffaf sekilde ayri
kolonda gosterilir, gizli bir donusum yok).

Sadece PARA/DEGER kolonlari cevrilir (fiyat, hacim, endeks kapanis degerleri).
Miktar (KG) ve adet kolonlari fiziksel oldugu icin cevrilmez.
"""
import glob
import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

KOK = Path(__file__).parent
GUNLUK_VERI_DIZINI = Path.home() / "Desktop" / "4-09-2026-turib" / "gunluk-bulten"  # agir veri, git disi (bkz PROJE_GECMISI.md)
HEDEF = Path("/Users/utkus/Desktop/turib_gunluk_bultenler_2023_2026_USD.xlsx")
KUR_CACHE = json.loads((KOK / "usd_kur_cache.json").read_text(encoding="utf-8"))


def sayiya_cevir(s):
    if s is None or s == "" or s == "-":
        return None
    s = str(s).strip().replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return s


# TL bazli PARA/DEGER kolonlari (USD'ye cevrilecek). Miktar[KG] ve Islem Adedi
# cevrilmez - bunlar fiziksel/sayisal, doviz kuruyla ilgisi yok.
PARA_ALANLARI = {
    "Önceki Kapanış", "En Düşük", "En Yüksek", "Kapanış",              # Endeks
    "Önceki Kapanış Fiyatı", "En Düşük Fiyat", "En Yüksek Fiyat",       # Normal Seans
    "Ağırlıklı Ortalama Fiyat", "Kapanış Fiyatı",
    "İşlem Hacmi [TL]",                                                  # Endeks + Normal Seans
}
# Cevrilmeyen ama tutulan sayisal alanlar
FIZIKSEL_ALANLAR = {"İşlem Adedi", "İşlem Miktarı [KG]"}


def basliklari_birlestir(satirlar: list[dict]) -> list[str]:
    gorulen, gorulen_set = [], set()
    for s in satirlar:
        for k in s.keys():
            if k not in gorulen_set:
                gorulen.append(k)
                gorulen_set.add(k)
    for sabit in ("tarih", "usd_kuru"):
        if sabit in gorulen:
            gorulen.remove(sabit)
    return ["tarih", "usd_kuru"] + gorulen


def sekme_yaz(wb, ad, satirlar):
    if not satirlar:
        return
    basliklar = basliklari_birlestir(satirlar)
    ws = wb.create_sheet(ad)
    ws.append(basliklar)
    for hucre in ws[1]:
        hucre.font = Font(bold=True)
    for satir in satirlar:
        ws.append([satir.get(b) for b in basliklar])
    for i, b in enumerate(basliklar, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, min(30, len(b) + 2))
    ws.freeze_panes = "A2"
    print(f"  [{ad}] {len(satirlar)} satir, {len(basliklar)} kolon yazildi")


def alan_isim_usd(alan: str) -> str:
    if alan in PARA_ALANLARI:
        return alan.replace("[TL]", "[USD]") if "[TL]" in alan else f"{alan} (USD)"
    return alan


def main():
    dosyalar = sorted(glob.glob(str(GUNLUK_VERI_DIZINI / "*.json")))
    print(f"{len(dosyalar)} gunluk bulten dosyasi bulundu")

    normal_seans_tum, endeks_tum, anlasmali_tum = [], [], []
    eksik_kur = set()

    for f in dosyalar:
        d = json.load(open(f, encoding="utf-8"))
        tarih = d["tarih"]
        kur_bilgi = KUR_CACHE.get(tarih)
        if kur_bilgi is None:
            eksik_kur.add(tarih)
            continue
        kur = kur_bilgi["kur"]

        for grup_adi, kaynak, hedef in (
            ("endeks", d["endeks"], endeks_tum),
            ("normal_seans", d["normal_seans"], normal_seans_tum),
        ):
            for s in kaynak:
                satir = {"tarih": tarih, "usd_kuru": kur}
                for k, v in s.items():
                    if k in PARA_ALANLARI:
                        sayi = sayiya_cevir(v)
                        satir[alan_isim_usd(k)] = round(sayi / kur, 4) if sayi is not None else None
                    elif k in FIZIKSEL_ALANLAR:
                        satir[k] = sayiya_cevir(v)
                    else:
                        satir[k] = v
                hedef.append(satir)

        for s in d["anlasmali"]:
            satir = {"tarih": tarih, "usd_kuru": kur}
            for k, v in s.items():
                satir[k] = sayiya_cevir(v) if k in FIZIKSEL_ALANLAR else v
            anlasmali_tum.append(satir)

    if eksik_kur:
        print(f"\nUYARI: {len(eksik_kur)} tarih icin USD kuru cache'te yok, bu gunler ATLANDI: "
              f"{sorted(eksik_kur)[:5]}{'...' if len(eksik_kur) > 5 else ''}")
        print("Once 'python3 tcmb_kur_cek.py' calistirip kur cache'ini tamamlayin.\n")

    print(f"Normal Seans: {len(normal_seans_tum)} satir, Endeks: {len(endeks_tum)} satir, "
          f"Anlasmali: {len(anlasmali_tum)} satir")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sekme_yaz(wb, "Normal Seans (USD)", normal_seans_tum)
    sekme_yaz(wb, "Endeks (USD)", endeks_tum)
    sekme_yaz(wb, "Anlaşmalı", anlasmali_tum)

    print("Kaydediliyor (buyuk dosya, biraz surebilir)...")
    wb.save(HEDEF)
    print(f"\nKaydedildi: {HEDEF}")


if __name__ == "__main__":
    main()
