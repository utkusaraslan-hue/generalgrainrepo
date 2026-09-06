"""TURIB'in 2023-2026 (bugune kadar) tum gunluk bultenlerini (Normal Seans,
Endeks, Anlasmali) tarih sirasina gore tek Excel dosyasinda birlestirir.

Baslik seti tum kayitlarin BIRLESIMI olarak cikarilir (4 yil boyunca kolon
seti degismis olabilir ihtimaline karsi), tarih her zaman ilk kolon.
"""
import glob
import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

KOK = Path(__file__).parent
GUNLUK_VERI_DIZINI = Path.home() / "Desktop" / "4-09-2026-turib" / "gunluk-bulten"  # agir veri, git disi (bkz PROJE_GECMISI.md)
HEDEF = Path("/Users/utkus/Desktop/turib_gunluk_bultenler_2023_2026.xlsx")


def sayiya_cevir(s):
    if s is None or s == "" or s == "-":
        return None
    s = str(s).strip().replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return s


SAYISAL_ALANLAR = {
    "Önceki Kapanış", "En Düşük", "En Yüksek", "Kapanış",
    "Önceki Kapanış Fiyatı", "En Düşük Fiyat", "En Yüksek Fiyat",
    "Ağırlıklı Ortalama Fiyat", "Kapanış Fiyatı", "İşlem Adedi",
    "İşlem Miktarı [KG]", "İşlem Hacmi [TL]",
}


def basliklari_birlestir(satirlar: list[dict]) -> list[str]:
    gorulen = []
    gorulen_set = set()
    for s in satirlar:
        for k in s.keys():
            if k not in gorulen_set:
                gorulen.append(k)
                gorulen_set.add(k)
    gorulen.remove("tarih")
    return ["tarih"] + gorulen


def sekme_yaz(wb, ad, satirlar):
    if not satirlar:
        return
    basliklar = basliklari_birlestir(satirlar)
    ws = wb.create_sheet(ad)
    ws.append(basliklar)
    for hucre in ws[1]:
        hucre.font = Font(bold=True)
    for satir in satirlar:
        ws.append([satir.get(b) for b in basliklar])
    for i, b in enumerate(basliklar, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, min(30, len(b) + 2))
    ws.freeze_panes = "A2"
    print(f"  [{ad}] {len(satirlar)} satir, {len(basliklar)} kolon yazildi")


def main():
    dosyalar = sorted(glob.glob(str(GUNLUK_VERI_DIZINI / "*.json")))
    print(f"{len(dosyalar)} gunluk bulten dosyasi bulundu")

    normal_seans_tum, endeks_tum, anlasmali_tum = [], [], []
    for f in dosyalar:
        d = json.load(open(f, encoding="utf-8"))
        tarih = d["tarih"]
        for s in d["endeks"]:
            satir = {"tarih": tarih, **s}
            for alan in SAYISAL_ALANLAR:
                if alan in satir:
                    satir[alan] = sayiya_cevir(satir[alan])
            endeks_tum.append(satir)
        for s in d["normal_seans"]:
            satir = {"tarih": tarih, **s}
            for alan in SAYISAL_ALANLAR:
                if alan in satir:
                    satir[alan] = sayiya_cevir(satir[alan])
            normal_seans_tum.append(satir)
        for s in d["anlasmali"]:
            anlasmali_tum.append({"tarih": tarih, **s})

    print(f"Normal Seans: {len(normal_seans_tum)} satir, Endeks: {len(endeks_tum)} satir, "
          f"Anlasmali: {len(anlasmali_tum)} satir")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sekme_yaz(wb, "Normal Seans", normal_seans_tum)
    sekme_yaz(wb, "Endeks", endeks_tum)
    sekme_yaz(wb, "Anlaşmalı", anlasmali_tum)

    print("Kaydediliyor (buyuk dosya, biraz surebilir)...")
    wb.save(HEDEF)
    print(f"\nKaydedildi: {HEDEF}")


if __name__ == "__main__":
    main()
