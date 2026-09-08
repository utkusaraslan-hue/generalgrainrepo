"""TURIB'in 2023'ten bugune TUM gunluk bultenlerini (Normal Seans + Endeks +
Anlasmali) iki Excel'de birlestirir - biri ham TL, biri o gunun TCMB USD
Satis kuruyla cevrilmis - ve HER İKİSİNE de "YoY Değişim (%)" kolonu ekler:
ayni (İl, İlçe, LİDAŞ Adı, Enstrüman Sınıfı) kombinasyonunun BİR ÖNCEKİ YILIN
AYNI TAKVİM GÜNÜNDEKİ kapanış fiyatına göre yüzde değişimi (Endeks sekmesinde
"Endeks Adı" bazinda). O gun TURIB kapali/yeni islem gormemisse (ya da bir
onceki yil ayni gun veri yoksa) hucre bos birakilir.

Cikti: ~/Desktop/turib-tarihi/turib_gunluk_bultenler_TL.xlsx
       ~/Desktop/turib-tarihi/turib_gunluk_bultenler_USD.xlsx

NOT: excel_birlestir.py / excel_birlestir_usd.py'nin YoY kolonlu genisletilmis
versiyonu - onlari degistirmez, ayri bir script'tir.
"""
import glob
import json
from datetime import date, timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

KOK = Path(__file__).parent
GUNLUK_VERI_DIZINI = Path.home() / "Desktop" / "4-09-2026-turib" / "gunluk-bulten"
HEDEF_DIZIN = Path.home() / "Desktop" / "turib-tarihi"
KUR_CACHE = json.loads((KOK / "usd_kur_cache.json").read_text(encoding="utf-8"))

PARA_ALANLARI = {
    "Önceki Kapanış", "En Düşük", "En Yüksek", "Kapanış",
    "Önceki Kapanış Fiyatı", "En Düşük Fiyat", "En Yüksek Fiyat",
    "Ağırlıklı Ortalama Fiyat", "Kapanış Fiyatı", "İşlem Hacmi [TL]",
}
FIZIKSEL_ALANLAR = {"İşlem Adedi", "İşlem Miktarı [KG]"}


def sayiya_cevir(s):
    if s is None or s == "" or s == "-":
        return None
    s = str(s).strip().replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return s


def onceki_yil_tarihi(tarih_iso: str) -> str:
    d = date.fromisoformat(tarih_iso)
    try:
        return d.replace(year=d.year - 1).isoformat()
    except ValueError:  # 29 Subat
        return (d.replace(year=d.year - 1, day=28)).isoformat()


def alan_isim_usd(alan: str) -> str:
    if alan in PARA_ALANLARI:
        return alan.replace("[TL]", "[USD]") if "[TL]" in alan else f"{alan} (USD)"
    return alan


def yoy_hesapla(gecmis: dict, anahtar, tarih_iso: str, deger) -> float | None:
    """gecmis: {anahtar: {tarih_iso: deger}}. Bir onceki yilin ayni gunundeki
    degeri bulup yuzde degisimi doner (yoksa None)."""
    if deger is None:
        return None
    onceki_tarih = onceki_yil_tarihi(tarih_iso)
    onceki_deger = gecmis.get(anahtar, {}).get(onceki_tarih)
    if onceki_deger is None or onceki_deger == 0:
        return None
    return round((deger - onceki_deger) / onceki_deger * 100, 2)


def basliklari_birlestir(satirlar: list[dict]) -> list[str]:
    gorulen, gorulen_set = [], set()
    for s in satirlar:
        for k in s.keys():
            if k not in gorulen_set:
                gorulen.append(k)
                gorulen_set.add(k)
    for sabit in ("tarih", "usd_kuru"):
        if sabit in gorulen:
            gorulen.remove(sabit)
    basliklar = [b for b in gorulen if b != "tarih" and b != "usd_kuru"]
    onek = ["tarih"] + (["usd_kuru"] if "usd_kuru" in gorulen_set else [])
    return onek + basliklar


def sekme_yaz(wb, ad, satirlar):
    if not satirlar:
        return
    basliklar = basliklari_birlestir(satirlar)
    ws = wb.create_sheet(ad)
    ws.append(basliklar)
    for hucre in ws[1]:
        hucre.font = Font(bold=True)
    for satir in satirlar:
        ws.append([satir.get(b) for b in basliklar])
    for i, b in enumerate(basliklar, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, min(30, len(b) + 2))
    ws.freeze_panes = "A2"
    print(f"  [{ad}] {len(satirlar)} satir, {len(basliklar)} kolon yazildi")


def main():
    dosyalar = sorted(glob.glob(str(GUNLUK_VERI_DIZINI / "*.json")))
    print(f"{len(dosyalar)} gunluk bulten dosyasi bulundu")

    HEDEF_DIZIN.mkdir(parents=True, exist_ok=True)

    ham = [json.load(open(f, encoding="utf-8")) for f in dosyalar]
    ham.sort(key=lambda d: d["tarih"])

    # --- 1) gecmis fiyat indeksleri (YoY icin, her iki para birimi ayri) ---
    ns_gecmis_tl, ns_gecmis_usd = {}, {}
    endeks_gecmis_tl, endeks_gecmis_usd = {}, {}
    eksik_kur = set()

    for d in ham:
        tarih = d["tarih"]
        kur_bilgi = KUR_CACHE.get(tarih)
        kur = kur_bilgi["kur"] if kur_bilgi else None
        if kur is None:
            eksik_kur.add(tarih)

        for s in d["normal_seans"]:
            anahtar = (s.get("İl"), s.get("İlçe"), s.get("LİDAŞ Adı"), s.get("Enstrüman Sınıfı"))
            kapanis = sayiya_cevir(s.get("Kapanış Fiyatı"))
            if kapanis is not None:
                ns_gecmis_tl.setdefault(anahtar, {})[tarih] = kapanis
                if kur:
                    ns_gecmis_usd.setdefault(anahtar, {})[tarih] = round(kapanis / kur, 4)

        for s in d["endeks"]:
            anahtar = s.get("Endeks Adı")
            kapanis = sayiya_cevir(s.get("Kapanış"))
            if kapanis is not None:
                endeks_gecmis_tl.setdefault(anahtar, {})[tarih] = kapanis
                if kur:
                    endeks_gecmis_usd.setdefault(anahtar, {})[tarih] = round(kapanis / kur, 4)

    if eksik_kur:
        print(f"UYARI: {len(eksik_kur)} tarih icin USD kuru yok, USD sekmelerinde bu gunler atlanacak "
              f"(once tcmb_kur_cek.py calistirin): {sorted(eksik_kur)[:5]}...")

    # --- 2) satirlari kur, YoY kolonuyla birlikte olustur ---
    def satirlari_olustur(usd_mi: bool):
        ns_tum, endeks_tum, anlasmali_tum = [], [], []
        for d in ham:
            tarih = d["tarih"]
            kur_bilgi = KUR_CACHE.get(tarih)
            kur = kur_bilgi["kur"] if kur_bilgi else None
            if usd_mi and kur is None:
                continue

            for s in d["normal_seans"]:
                satir = {"tarih": tarih}
                if usd_mi:
                    satir["usd_kuru"] = kur
                anahtar = (s.get("İl"), s.get("İlçe"), s.get("LİDAŞ Adı"), s.get("Enstrüman Sınıfı"))
                kapanis_tl = sayiya_cevir(s.get("Kapanış Fiyatı"))
                for k, v in s.items():
                    if k in PARA_ALANLARI:
                        sayi = sayiya_cevir(v)
                        if usd_mi:
                            satir[alan_isim_usd(k)] = round(sayi / kur, 4) if sayi is not None else None
                        else:
                            satir[k] = sayi
                    elif k in FIZIKSEL_ALANLAR:
                        satir[k] = sayiya_cevir(v)
                    else:
                        satir[k] = v
                    if k == "Kapanış Fiyatı":
                        if usd_mi:
                            deger = round(kapanis_tl / kur, 4) if kapanis_tl is not None else None
                            satir["YoY Değişim (%)"] = yoy_hesapla(ns_gecmis_usd, anahtar, tarih, deger)
                        else:
                            satir["YoY Değişim (%)"] = yoy_hesapla(ns_gecmis_tl, anahtar, tarih, kapanis_tl)
                ns_tum.append(satir)

            for s in d["endeks"]:
                satir = {"tarih": tarih}
                if usd_mi:
                    satir["usd_kuru"] = kur
                anahtar = s.get("Endeks Adı")
                kapanis_tl = sayiya_cevir(s.get("Kapanış"))
                for k, v in s.items():
                    if k in PARA_ALANLARI:
                        sayi = sayiya_cevir(v)
                        if usd_mi:
                            satir[alan_isim_usd(k)] = round(sayi / kur, 4) if sayi is not None else None
                        else:
                            satir[k] = sayi
                    elif k in FIZIKSEL_ALANLAR:
                        satir[k] = sayiya_cevir(v)
                    else:
                        satir[k] = v
                    if k == "Kapanış":
                        if usd_mi:
                            deger = round(kapanis_tl / kur, 4) if kapanis_tl is not None else None
                            satir["YoY Değişim (%)"] = yoy_hesapla(endeks_gecmis_usd, anahtar, tarih, deger)
                        else:
                            satir["YoY Değişim (%)"] = yoy_hesapla(endeks_gecmis_tl, anahtar, tarih, kapanis_tl)
                endeks_tum.append(satir)

            for s in d["anlasmali"]:
                satir = {"tarih": tarih}
                if usd_mi:
                    satir["usd_kuru"] = kur
                for k, v in s.items():
                    satir[k] = sayiya_cevir(v) if k in FIZIKSEL_ALANLAR else v
                anlasmali_tum.append(satir)

        return ns_tum, endeks_tum, anlasmali_tum

    for usd_mi, dosya_adi, sekme_onek in (
        (False, "turib_gunluk_bultenler_TL.xlsx", ""),
        (True, "turib_gunluk_bultenler_USD.xlsx", " (USD)"),
    ):
        print(f"\n=== {'USD' if usd_mi else 'TL'} versiyonu uretiliyor ===")
        ns_tum, endeks_tum, anlasmali_tum = satirlari_olustur(usd_mi)
        print(f"Normal Seans: {len(ns_tum)} satir, Endeks: {len(endeks_tum)} satir, "
              f"Anlaşmalı: {len(anlasmali_tum)} satir")
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        sekme_yaz(wb, f"Normal Seans{sekme_onek}", ns_tum)
        sekme_yaz(wb, f"Endeks{sekme_onek}", endeks_tum)
        sekme_yaz(wb, "Anlaşmalı", anlasmali_tum)
        hedef = HEDEF_DIZIN / dosya_adi
        print("Kaydediliyor (buyuk dosya, biraz surebilir)...")
        wb.save(hedef)
        print(f"Kaydedildi: {hedef}")


if __name__ == "__main__":
    main()
