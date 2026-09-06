# -*- coding: utf-8 -*-
"""TURIB sitesinden cikarilan TUM sayisal ucret/tarife verilerini tek Excel'de toplar."""
import csv
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

KOK = Path(__file__).parent
HEDEF = KOK / "turib_ucret_tarifeleri_ve_veriler.xlsx"

wb = openpyxl.Workbook()
wb.remove(wb.active)


def sekme(ad, basliklar, satirlar, kaynak_url=None, not_metni=None):
    ws = wb.create_sheet(ad[:31])
    satir_no = 1
    if not_metni:
        ws.cell(row=satir_no, column=1, value=not_metni).font = Font(italic=True, color="666666")
        satir_no += 2
    if kaynak_url:
        ws.cell(row=satir_no, column=1, value=f"Kaynak: {kaynak_url}").font = Font(italic=True, size=9, color="888888")
        satir_no += 2
    baslik_satiri = satir_no
    for i, b in enumerate(basliklar, start=1):
        c = ws.cell(row=baslik_satiri, column=i, value=b)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = openpyxl.styles.PatternFill("solid", fgColor="34495E")
    for r, satir in enumerate(satirlar, start=baslik_satiri + 1):
        for c_i, deger in enumerate(satir, start=1):
            ws.cell(row=r, column=c_i, value=deger)
    for i, b in enumerate(basliklar, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(45, len(str(b)) + 4))
    ws.freeze_panes = ws.cell(row=baslik_satiri + 1, column=1).coordinate
    return ws


# ============ 1) Depolama Ucret Tarifesi (Hububat/Baklagil/Yagli Tohum) ============
hub_basliklar = ["No", "Ürün", "28.06.2022 (TL/Ton/Gün)", "23.01.2023", "31.01.2024", "31.01.2025", "30.01.2026"]
hub_satirlar = [
    [1, "Buğday", 0.6, 1.08, 1.66, 2.27, 2.85],
    [2, "Çavdar", 0.6, 1.08, 1.66, 2.27, 2.85],
    [3, "Tritikale", 0.6, 1.08, 1.66, 2.27, 2.85],
    [4, "Mısır", 0.63, 1.14, 1.76, 2.4, 3.01],
    [5, "Arpa", 0.67, 1.21, 1.86, 2.54, 3.19],
    [6, "Çeltik (Dökme)", 0.67, 1.21, 1.86, 2.54, 3.19],
    [7, "Çeltik (Çuvallı)", 0.87, 1.57, 2.42, 3.3, 4.14],
    [8, "Yulaf", 0.67, 1.21, 1.86, 2.54, 3.19],
    [9, "Kanola (Kolza)", 0.67, 1.21, 1.86, 2.54, 3.19],
    [10, "Nohut", 0.83, 1.5, 2.31, 3.15, 3.95],
    [11, "Mercimek", 0.83, 1.5, 2.31, 3.15, 3.95],
    [12, "Fasulye", 0.83, 1.5, 2.31, 3.15, 3.95],
    [13, "Soya Fasulyesi", 0.83, 1.5, 2.31, 3.15, 3.95],
    [14, "Ayçiçeği", 1.03, 1.86, 2.87, 3.92, 4.92],
]
sekme("Depolama Ucreti-Hububat", hub_basliklar, hub_satirlar,
      "https://www.turib.com.tr/lisansli-depoculuk-depolama-ucret-tarifesi-uygulamasi/",
      "KDV HARIC fiyatlardir. Birim: TL/Ton/Gun.")

sekme("Depolama Ucreti-Diger Urunler",
      ["Ürün", "Birim", "Tarih 1", "Değer 1", "Tarih 2", "Değer 2", "Tarih 3", "Değer 3", "Tarih 4", "Değer 4", "Tarih 5", "Değer 5", "KDV"],
      [
        ["Antep Fıstığı", "TL/Ton/Gün", "12.09.2022", 4.42, "31.01.2024", 7.3, "31.01.2025", 9.96, "30.01.2026", 12.88, "", "", "Hariç"],
        ["Pamuk (Rollergin Prese)", "TL/Kg/Ay", "20.09.2022", 0.082, "31.01.2023", 0.135, "31.01.2024", 0.210, "31.01.2025", 0.287, "30.01.2026", 0.371, "Hariç"],
        ["Pamuk (Sawgin Prese)", "TL/Kg/Ay", "20.09.2022", 0.082, "31.01.2023", 0.135, "31.01.2024", 0.210, "31.01.2025", 0.287, "30.01.2026", 0.371, "Hariç"],
        ["Zeytin (Mislen Depolama)", "TL/Ton/Ay", "18.11.2022", 50, "31.01.2024", 100, "31.01.2025", 136.45, "30.01.2026", 176.4, "", "", "DAHIL"],
        ["Kuru Kayısı", "TL/Ton/Gün", "30.12.2022", 4.35, "31.01.2024", 6.7, "31.01.2025", 9.14, "30.01.2026", 11.82, "", "", "Hariç"],
        ["Fındık", "TL/Ton/Gün", "17.08.2022", 1.5, "31.01.2023", 2.33, "31.01.2024", 4.00, "31.01.2025", 5.00, "30.01.2026", 6.46, "DAHIL"],
        ["Peynir (Süt Ürünleri)", "TL/Ton/Ay", "31.01.2024", 800, "31.01.2025", 1091.60, "", "", "", "", "", "", "Hariç"],
        ["Çekirdeksiz Kuru Üzüm", "TL/Ton/Gün", "28.10.2025", 14.65, "", "", "", "", "", "", "", "", "Hariç"],
      ],
      "https://www.turib.com.tr/lisansli-depoculuk-depolama-ucret-tarifesi-uygulamasi/")

# ============ 2) Fire Tarifesi ============
sekme("Fire Tarifesi",
      ["Ürün Grubu", "Depolama Süresi", "Fire Oranı", "Günlük Fire Katsayısı", "Formül (Tek Dönem)"],
      [
        ["Hububat, Baklagiller, Yağlı Tohumlar", "365 güne kadar", "0,0035", "0,000009589041",
         "ELÜS Miktarı x İşlem Fiyatı x Fire Oranı x Fire Gün / 365"],
        ["Fındık", "730 güne kadar", "%1", "0,00001369863",
         "ELÜS Miktarı x İşlem Fiyatı x Fire Oranı x Fire Gün / 730"],
      ],
      "https://www.turib.com.tr/lisansli-depoculuk-fire-tarifesi-uygulamasi/",
      "Fire Tarifesi Başlangıç Tarihi: 02.09.2022. Fire miktarı = Günlük Fire Katsayısı x Fire Gün x Ürün Miktarı. "
      "Fire sınırı aşılırsa: Borsada oluşan fiyatın (son 5 işlem günü ağırlıklı ort., yoksa son 30 gün) %5 üzerinde bedel ödenir, 7 iş günü içinde teslim.")

# ============ 3) Depolama Donemleri (ne kadar tutulabilir) ============
sekme("Depolama Donemleri (Max Sure)",
      ["Ürün Grubu", "Tür", "Maksimum Saklama Süresi", "Dönem Başlangıç", "Dönem Bitiş"],
      [
        ["Hububat/Baklagil/Yağlı Tohum", "Buğday", "24 Ay", "15 Mayıs", "14 Mayıs"],
        ["Hububat/Baklagil/Yağlı Tohum", "Arpa", "24 Ay", "15 Mayıs", "14 Mayıs"],
        ["Hububat/Baklagil/Yağlı Tohum", "Çavdar", "24 Ay", "1 Haziran", "31 Mayıs"],
        ["Hububat/Baklagil/Yağlı Tohum", "Yulaf", "24 Ay", "1 Haziran", "31 Mayıs"],
        ["Hububat/Baklagil/Yağlı Tohum", "Mısır", "24 Ay", "1 Ağustos", "31 Temmuz"],
        ["Hububat/Baklagil/Yağlı Tohum", "Çeltik", "24 Ay", "1 Eylül", "31 Ağustos"],
        ["Hububat/Baklagil/Yağlı Tohum", "Pirinç", "24 Ay", "1 Eylül", "31 Ağustos"],
        ["Hububat/Baklagil/Yağlı Tohum", "Tritikale", "24 Ay", "15 Haziran", "14 Haziran"],
        ["Hububat/Baklagil/Yağlı Tohum", "Mercimek", "24 Ay", "1 Mayıs", "30 Nisan"],
        ["Hububat/Baklagil/Yağlı Tohum", "Nohut", "24 Ay", "15 Haziran", "14 Haziran"],
        ["Hububat/Baklagil/Yağlı Tohum", "Fasulye", "24 Ay", "1 Ağustos", "31 Temmuz"],
        ["Hububat/Baklagil/Yağlı Tohum", "Bakla", "24 Ay", "15 Temmuz", "14 Temmuz"],
        ["Hububat/Baklagil/Yağlı Tohum", "Ayçiçeği", "24 Ay", "15 Temmuz", "14 Temmuz"],
        ["Hububat/Baklagil/Yağlı Tohum", "Keten Tohumu", "24 Ay", "15 Temmuz", "14 Temmuz"],
        ["Hububat/Baklagil/Yağlı Tohum", "Soya Fasulyesi", "24 Ay", "15 Temmuz", "14 Temmuz"],
        ["Hububat/Baklagil/Yağlı Tohum", "Kolza", "24 Ay", "15 Mayıs", "14 Mayıs"],
        ["Pamuk", "Pamuk", "2 Yıl", "1 Eylül", "31 Ağustos"],
        ["Fındık", "Fındık", "2 Yıl", "15 Ağustos", "14 Ağustos"],
        ["Zeytin", "Siyah ve rengi dönük", "2 Yıl", "1 Kasım", "28 Şubat"],
        ["Zeytin", "Yeşil", "2 Yıl", "1 Eylül", "31 Ağustos"],
        ["Kuru Kayısı", "Kükürtsüz", "1 Yıl", "1 Ağustos", "31 Temmuz"],
        ["Kuru Kayısı", "Kükürtlü 2000-2500ppm", "9 Ay", "1 Ağustos", "30 Nisan"],
        ["Kuru Kayısı", "Kükürtlü 2501-3000ppm", "1 Yıl", "1 Ağustos", "31 Temmuz"],
        ["Kuru Kayısı", "Kükürtlü 3001-3500ppm", "15 Ay", "1 Ağustos", "31 Ekim"],
        ["Kuru Kayısı", "Kükürtlü 3501-üzeri", "18 Ay", "1 Ağustos", "31 Ocak"],
        ["Antep Fıstığı", "Ben Kavlak", "2 Yıl", "15 Ağustos", "14 Ağustos"],
        ["Antep Fıstığı", "Boz Kavlak", "2 Yıl", "1 Ağustos", "31 Temmuz"],
        ["Antep Fıstığı", "Kırmızı Kabuklu", "2 Yıl", "1 Eylül", "31 Ağustos"],
        ["Antep Fıstığı", "Meverdi Kavlak", "2 Yıl", "5 Ağustos", "4 Ağustos"],
        ["Süt Ürünleri", "Peynir", "12 Ay", "1 Ocak", "31 Aralık"],
        ["Kuru Üzüm", "Çekirdeksiz Kuru Üzüm", "12 Ay", "15 Ağustos", "14 Ağustos"],
      ],
      "https://www.turib.com.tr/urunlerin-depolama-donemleri/",
      "NOT: 2026 hasat donemi oncesi teslim edilen urunler icin azami sure TMO urunlerinde 36 ay, diger mudilerde 18 aydir.")

# ============ 4) ELUS Islem Ucretleri ============
sekme("ELUS Islem Ucretleri",
      ["Ücret Türü", "Oran", "Ondalık", "Açıklama"],
      [
        ["Borsa Tescil Ücreti", "Binde 0,25 (On binde 2,5)", 0.00025, "Her ELÜS işleminde, takas anında, işlem taraflarının HER BİRİNDEN, işlem tutarı üzerinden (+KDV)"],
        ["Borsa Hizmet Ücreti", "Binde 1,25", 0.00125, "Her ELÜS işleminde, takas anında, işlem taraflarının HER BİRİNDEN, işlem tutarı üzerinden (+KDV)"],
        ["Lisanslı Depoculuk Tazmin Fonu Payı", "Binde 0,25 (On binde 2,5)", 0.00025, "Her ELÜS işleminde, takas anında, işlem taraflarının HER BİRİNDEN, işlem tutarı üzerinden (KDV YOK)"],
      ],
      "https://www.turib.com.tr/turib-elus-piyasasi-ve-fon-payi-hesaplama/",
      "ÖRNEK (1.000.000 TL işlem): Tescil 250 TL+50 KDV=300; Tazmin Fonu 250 TL (KDV yok); Hizmet 1.250 TL+250 KDV=1.500 -> TOPLAM 2.050 TL (tek taraf için)")

# ============ 5) Kantar / Seviye Olcum Sistemi Ucreti ============
sekme("Kantar-Seviye Olcum Ucreti",
      ["Kalem", "Değer", "Açıklama"],
      [
        ["Birim ücret (vergi dahil maliyet katsayısı)", "5,31 TL/Ton", "Lisanslı Depo Lisans Kapasitesi x 5,31 TL/Ton = Ödeme Tutarı"],
        ["Toplam sistem yazılım+kurulum maliyeti (KDV hariç)", "61.030.800 TL", "Tüm sektöre dağıtılan toplam maliyet"],
        ["Toplam sistem yazılım+kurulum maliyeti (KDV dahil, %20)", "73.236.960 TL", ""],
        ["Son ödeme tarihi (2026 duyurusu)", "27.02.2026 mesai bitimi", "Bu tarihten sonra faaliyet izni alacaklar icin TUFE farki yansitilir"],
      ],
      "https://www.turib.com.tr/lisansli-depo-bilgi-sistemi-kapsaminda-lisansli-depoculuk-seviye-olcum-sistemi-ucret-tarifesi-2/ (+ PDF: Lisansli_Depoculuk_Seviye_Olcum_Sistemi_Ucret_Tarifesi.pdf)",
      "Sadece Hububat/Baklagil/Yagli Tohum lisansli depolari icin (2026 ilk modul). Gecikme zammi: aylik TUFE orani gun bazinda uygulanir.")

# ============ 6) Veri Yayim Hizmeti Ucret Tarifesi ============
sekme("Veri Yayim Ucret Tarifesi (2026)",
      ["Paket", "Sabit Dağıtım Bedeli (TL/ay+KDV)", "Değişken Dağıtım Bedeli", "100 Kullanıcı (TL/ay)", "500 Kullanıcı", "1000 Kullanıcı", "Aşım Bedeli(100k)", "Aşım(500k)", "Aşım(1000k)"],
      [
        ["Paket 0", 5395.20, "Ücretsiz", "-", "-", "-", "-", "-", "-"],
        ["Paket 1", 2697.60, "Ücretsiz", "-", "-", "-", "-", "-", "-"],
        ["Paket 2", 8092.80, "31.12.2026'ya kadar alınmaz", 6744, 30348, 53952, 67.44, 60.70, 53.95],
        ["Paket 3", 13488.00, "31.12.2026'ya kadar alınmaz", 13488, 60696, 107904, 134.88, 121.39, 107.90],
        ["TV/Web Paketi", 5395.20, "Ücretsiz (kanal/site başına)", "-", "-", "-", "-", "-", "-"],
      ],
      "https://www.turib.com.tr/turib-veri-yayim-hizmeti-ucret-tarifesi/",
      "01.01.2027'den itibaren TUFE yillik degisim oranina gore artirilir.")

# ============ 7) Borsa Uyelik Ucretleri ============
sekme("Borsa Uyelik Ucretleri (2026)",
      ["Kalem", "Tutar (KDV hariç)", "Ödeme Yöntemi", "Bakanlık Onay Tarihi"],
      [
        ["ELÜS Piyasası Üyelik Kayıt Ücreti", "812.780 TL", "Üyeliğe ilk girişte", "11.10.2024"],
        ["ELÜS Piyasası Yıllık Üyelik Aidatı", "541.853 TL", "Her yıl Ocak ayının 15. iş gününe kadar (başvuru yılı hariç)", "11.10.2024"],
        ["ELÜS Piyasası Üyelik Teminatı", "541.853 TL", "2026 yılı sonuna kadar alınmayacak", "-"],
      ],
      "https://www.turib.com.tr/borsaya-uyelik/",
      "Her yil Ocak ayinda VUK mük. 298/B fikrasina gore yeniden degerleme oraninda artirilir.")

# ============ 8) Ciftci Destekleri (Kira/Nakliye/Analiz) ============
sekme("Ciftci Destekleri",
      ["Ürün Grubu", "Kira Desteği", "Nakliye Desteği (TL/Ton)", "Nakliye Üst Sınır", "Analiz Desteği"],
      [
        ["Buğday, Arpa, Yulaf, Çavdar, Mısır, Çeltik, Pirinç", "Bakanlık kira tarifesinin %75'i (2025-2027)", 100, "30 ton", "Bakanlık analiz tarifesinin %50'si"],
        ["Mercimek, Nohut, Fasulye, Soya Fasulyesi, Ayçiçeği, Pamuk, Fındık, Antep Fıstığı, Kuru Kayısı/Üzüm, Zeytin/Zeytinyağı",
         "Bakanlık kira tarifesinin %75'i (2025-2027)", 300, "10 ton", "Bakanlık analiz tarifesinin %50'si"],
      ],
      "https://www.turib.com.tr/lisansli-depoculuk-depolama-ucret-tarifesi-uygulamasi/ + https://www.turib.com.tr/ciftcilerimize-verilen-kira-nakliye-ve-analiz-desteklerinde-onemli-yenilikler-yapildi/",
      "TARIHCE: 2024 duyurusunda kira destegi %60'ti (5,5 kata varan artis oncesi seviye), 2025-2027 icin %75'e yukseltildi. "
      "Kira destegi en fazla 6 ay sure ile odenir, ÇKS kayitli ureticiler/birlikler/kooperatifler icin gecerli. "
      "Vergi istisnasi: 2028 sonuna kadar gelir/kurumlar vergisi muafiyeti + damga vergisi/KDV istisnasi. "
      "Ziraat Bankasi sifir faizli kredi: ELUS tutarinin %75'ine kadar, azami 9 ay vade.")

# ============ 9) ELUS ISIN Kod Ucreti ============
sekme("ELUS ISIN Kod Ucreti",
      ["Ücret Adı", "Eski (12.02.2020 öncesi)", "Yeni (12.02.2020'den itibaren)"],
      [["Elektronik Ürün Sertifikası ISIN kodu tahsis bedeli", "100 TL + BSMV", "50 TL + BSMV"]],
      "https://www.turib.com.tr/elus-isin-kodu-tahsis-ucretinde-indirim/")

print("Sekmeler tamamlandi, depo listesi ekleniyor...")

# ============ 10) Depo ISIN Kod Listesi (368 depo) ============
depo_satirlari = []
with open(KOK / "depo_isin_kod_listesi.csv", encoding="utf-8") as f:
    r = csv.reader(f)
    next(r)
    for row in r:
        depo_satirlari.append(row)

sekme("Lisansli Depo Listesi (368)",
      ["Sıra No", "İhraç Kodu", "LİDAŞ Ünvanı", "İl", "İlçe/Lokasyon"],
      depo_satirlari,
      "https://www.turib.com.tr/lisansli-depo-isletmesi-isin-kodlari/")

wb.save(HEDEF)
print(f"Kaydedildi: {HEDEF}")
